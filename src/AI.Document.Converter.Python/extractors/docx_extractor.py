"""DOCX extraction (FR-007) via python-docx.

python-docx exposes document.paragraphs and document.tables as two separate,
unordered collections, which would lose the original reading order of a document
that interleaves paragraphs and tables. _iter_block_items below is the standard
recipe for walking the document body in its actual XML order instead.
"""

import os
import re
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table
from docx.text.paragraph import Paragraph

from extractors import model
from extractors.common import (
    ExtractionError,
    check_file_accessible,
    check_not_encrypted_ooxml,
    converted_date_iso,
    file_created_date_iso,
    image_placeholder_block,
    normalize_optional_text,
    unextractable_text_block,
)


def _iter_block_items(document):
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, document)
        elif isinstance(child, CT_Tbl):
            yield Table(child, document)


def _heading_level(style_name):
    match = re.search(r"(\d+)", style_name or "")
    return int(match.group(1)) if match else 1


def _paragraph_text_with_inline_links(paragraph):
    # paragraph.text (python-docx's own property) already walks into
    # w:hyperlink elements for their visible text, so a naive "plain text plus
    # a separate LinkBlock per hyperlink" would duplicate the link text in the
    # output. Markdown represents an inline link as `[text](url)` inline in the
    # running prose anyway, so building that directly here - rather than a
    # standalone LinkBlock - is both correct and the more natural fit; LinkBlock
    # remains available for contexts where a link isn't embedded in prose.
    parts = []
    for child in paragraph._p:
        if child.tag == qn("w:r"):
            parts.append("".join(node.text or "" for node in child.iter(qn("w:t"))))
        elif child.tag == qn("w:hyperlink"):
            text = "".join(node.text or "" for node in child.iter(qn("w:t"))).strip()
            r_id = child.get(qn("r:id"))
            url = None
            if r_id:
                try:
                    url = paragraph.part.rels[r_id].target_ref
                except KeyError:
                    url = None
            parts.append(f"[{text}]({url})" if text and url else text)
    return "".join(parts).strip()


def _run_has_image(run_element):
    # w:drawing is the modern (DrawingML) inline-image element python-docx's
    # own add_picture() produces; w:pict is the legacy VML form still seen
    # in older/converted documents. Either one means "this run is an image."
    return run_element.find(qn("w:drawing")) is not None or run_element.find(qn("w:pict")) is not None


def _paragraph_image_count(paragraph):
    # FR-039/AC-027: an inline image lives inside a w:r (or a w:r nested
    # inside a w:hyperlink) as a sibling of w:t, not as a separate top-level
    # body element - _iter_block_items only sees the containing Paragraph,
    # so without this the image is silently dropped rather than marked.
    count = 0
    for child in paragraph._p:
        if child.tag == qn("w:r") and _run_has_image(child):
            count += 1
        elif child.tag == qn("w:hyperlink"):
            count += sum(1 for run in child if run.tag == qn("w:r") and _run_has_image(run))
    return count


def extract(file_path):
    check_file_accessible(file_path)
    check_not_encrypted_ooxml(file_path)

    try:
        document = Document(file_path)
    except (PackageNotFoundError, BadZipFile) as ex:
        raise ExtractionError(
            "corruptedDocument", f"'{os.path.basename(file_path)}' could not be opened: {ex}"
        )

    sections = []
    current_section = {"heading": None, "headingLevel": None, "blocks": [], "location": None}
    pending_list = None

    def flush_list():
        nonlocal pending_list
        if pending_list is not None and pending_list["items"]:
            current_section["blocks"].append(pending_list)
        pending_list = None

    def start_new_section(heading_text, level):
        nonlocal current_section
        flush_list()
        if current_section["heading"] or current_section["blocks"]:
            sections.append(current_section)
        current_section = {"heading": heading_text, "headingLevel": level, "blocks": [], "location": None}

    for item in _iter_block_items(document):
        if isinstance(item, Paragraph):
            style_name = item.style.name if item.style else ""
            text = _paragraph_text_with_inline_links(item)
            image_count = _paragraph_image_count(item)

            if style_name.startswith("Heading") or style_name == "Title":
                start_new_section(text, _heading_level(style_name))
                for _ in range(image_count):
                    current_section["blocks"].append(image_placeholder_block())
                continue

            if "List" in style_name:
                is_ordered = "Number" in style_name
                if pending_list is None or pending_list["isOrdered"] != is_ordered:
                    flush_list()
                    pending_list = {"type": "list", "isOrdered": is_ordered, "items": []}
                if text:
                    pending_list["items"].append(text)
                for _ in range(image_count):
                    current_section["blocks"].append(image_placeholder_block())
                continue

            flush_list()

            if text:
                current_section["blocks"].append({"type": "paragraph", "text": text})

            for _ in range(image_count):
                current_section["blocks"].append(image_placeholder_block())

        elif isinstance(item, Table):
            flush_list()
            rows = [[cell.text for cell in row.cells] for row in item.rows]
            if rows:
                current_section["blocks"].append(
                    {"type": "table", "headers": rows[0], "rows": rows[1:]}
                )

    flush_list()
    if current_section["heading"] or current_section["blocks"]:
        sections.append(current_section)

    if not sections:
        sections.append(
            {
                "heading": None,
                "headingLevel": None,
                "blocks": [unextractable_text_block("Document has no extractable content.")],
                "location": None,
            }
        )

    core_properties = document.core_properties

    warnings = []

    images_omitted = sum(
        1
        for section in sections
        for block in section["blocks"]
        if block.get("type") == "imagePlaceholder"
    )
    if images_omitted:
        warnings.append(
            model.warning(
                model.IMAGE_OMITTED,
                model.INFO,
                (
                    f"{images_omitted} embedded image(s) were not extracted. A placeholder marks "
                    f"each position in the output."
                ),
                details={"imageCount": images_omitted},
            )
        )

    has_text = any(
        block.get("type") != "unextractableText"
        for section in sections
        for block in section["blocks"]
    )
    if not has_text:
        warnings.append(
            model.warning(
                model.NO_EXTRACTABLE_TEXT,
                model.ERROR,
                "No text content could be extracted from this document.",
            )
        )

    metadata = {
        "sourceFilePath": file_path,
        "fileType": "docx",
        "createdDate": file_created_date_iso(file_path),
        "convertedDate": converted_date_iso(),
        # No reliable layout-independent page count exists for DOCX, so the
        # field is omitted rather than invented (SaaS SR-INT / audit C-05).
        "pageCount": None,
        "slideCount": None,
        "sheetCount": None,
        "author": normalize_optional_text(core_properties.author),
    }

    return model.document(metadata, sections, warnings)
