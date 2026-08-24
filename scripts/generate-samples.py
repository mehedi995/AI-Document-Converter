"""Generates the synthetic, non-confidential sample files in samples/
(CLAUDE.md Section 34; docs/16-TEST-STRATEGY.md Section 3).

Run from the repo root: python scripts/generate-samples.py
Requires the packages in src/AI.Document.Converter.Python/requirements.txt.
"""

import base64
import os

import pymupdf
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.opc.constants import RELATIONSHIP_TYPE
import openpyxl
from pptx import Presentation

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAMPLES_DIR = os.path.join(REPO_ROOT, "samples")

# A real (if trivial) 1x1 PNG - AC-027 needs a genuine embedded image, not
# just bytes claiming to be one, since it's testing that each extractor's
# own image-detection logic actually fires on real embedded-image markup.
_TINY_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def add_hyperlink(paragraph, url, text):
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    run.append(run_properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)

    paragraph._p.append(hyperlink)


def generate_txt():
    path = os.path.join(SAMPLES_DIR, "sample.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(
            "Sample Text Document\n\n"
            "This is a plain text file used as a fixture for the AI Document "
            "Converter (FR-010).\n\n"
            "It has multiple paragraphs and no special structure, since TXT\n"
            "files are read as plain text without further parsing.\n"
        )
    print(f"Wrote {path}")


def generate_docx():
    path = os.path.join(SAMPLES_DIR, "sample.docx")
    document = docx.Document()

    document.add_heading("Sample Document", level=1)
    document.add_paragraph(
        "This is a sample DOCX fixture exercising headings, paragraphs, "
        "lists, tables, and hyperlinks (FR-007)."
    )

    document.add_heading("Key Points", level=2)
    document.add_paragraph("First point", style="List Bullet")
    document.add_paragraph("Second point", style="List Bullet")
    document.add_paragraph("Third point", style="List Bullet")

    document.add_heading("Reference Table", level=2)
    table = document.add_table(rows=3, cols=2)
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    header_cells[0].text = "Item"
    header_cells[1].text = "Value"
    table.rows[1].cells[0].text = "Alpha"
    table.rows[1].cells[1].text = "100"
    table.rows[2].cells[0].text = "Beta"
    table.rows[2].cells[1].text = "200"

    link_paragraph = document.add_paragraph("See also: ")
    add_hyperlink(link_paragraph, "https://example.com/reference", "example reference")

    document.save(path)
    print(f"Wrote {path}")


def generate_xlsx():
    path = os.path.join(SAMPLES_DIR, "sample.xlsx")
    workbook = openpyxl.Workbook()

    data_sheet = workbook.active
    data_sheet.title = "Data"
    data_sheet.append(["Item", "Quantity", "Unit Price", "Total"])
    data_sheet.append(["Widget", 10, 2.5, "=B2*C2"])
    data_sheet.append(["Gadget", 5, 9.99, "=B3*C3"])
    data_sheet.merge_cells("A5:B5")
    data_sheet["A5"] = "Merged note spanning two columns"

    # FR-041: a sheet large enough to exceed the extractor's row-cap threshold,
    # to exercise the summarization path with a still-tiny fixture file.
    large_sheet = workbook.create_sheet("LargeData")
    large_sheet.append(["Row Number", "Value"])
    for i in range(1, 251):
        large_sheet.append([i, f"value-{i}"])

    workbook.save(path)
    print(f"Wrote {path}")


def generate_pptx():
    path = os.path.join(SAMPLES_DIR, "sample.pptx")
    presentation = Presentation()

    title_slide_layout = presentation.slide_layouts[1]
    slide = presentation.slides.add_slide(title_slide_layout)
    slide.shapes.title.text = "Sample Presentation"
    body = slide.placeholders[1]
    text_frame = body.text_frame
    text_frame.text = "First bullet point"
    paragraph = text_frame.add_paragraph()
    paragraph.text = "Second bullet point"

    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "Speaker notes for the sample slide."

    second_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    second_slide.shapes.title.text = "Second Slide"
    second_slide.placeholders[1].text_frame.text = "Content on the second slide."

    presentation.save(path)
    print(f"Wrote {path}")


def generate_pdf():
    path = os.path.join(SAMPLES_DIR, "sample.pdf")
    document = pymupdf.open()

    page = document.new_page()
    page.insert_text((72, 72), "Sample PDF Document", fontsize=20, fontname="helv")
    page.insert_text(
        (72, 110),
        "This page exercises headings, body text, and a table (FR-006).",
        fontsize=11,
    )

    # A grid-lined table pymupdf's find_tables() can detect.
    table_top = 150
    row_height = 20
    col_widths = [100, 100]
    rows = [["Item", "Value"], ["Alpha", "100"], ["Beta", "200"]]

    for row_index, row in enumerate(rows):
        y0 = table_top + row_index * row_height
        y1 = y0 + row_height
        x = 72
        for col_index, cell_text in enumerate(row):
            width = col_widths[col_index]
            rect = pymupdf.Rect(x, y0, x + width, y1)
            page.draw_rect(rect, color=(0, 0, 0), width=0.75)
            page.insert_textbox(rect, cell_text, fontsize=10, align=1)
            x += width

    second_page = document.new_page()
    second_page.insert_text(
        (72, 72), "This is page two of the sample PDF, for page references.", fontsize=11
    )

    # FR-043: a page with no extractable text at all (e.g., a scanned page).
    document.new_page()

    document.save(path)
    print(f"Wrote {path}")


def generate_image_pdf():
    # AC-027/FR-039: a dedicated fixture, separate from sample.pdf, so this
    # scenario is isolated and doesn't force renumbering sample.pdf's
    # existing section/page assertions.
    path = os.path.join(SAMPLES_DIR, "image-sample.pdf")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "Text before the embedded image.", fontsize=11)
    page.insert_image(pymupdf.Rect(72, 100, 172, 200), stream=_TINY_PNG_BYTES)
    document.save(path)
    print(f"Wrote {path}")


def generate_image_docx():
    path = os.path.join(SAMPLES_DIR, "image-sample.docx")
    document = docx.Document()
    document.add_paragraph("Text before the embedded image.")
    image_paragraph = document.add_paragraph()
    png_path = os.path.join(SAMPLES_DIR, "_tmp-image-fixture.png")
    with open(png_path, "wb") as f:
        f.write(_TINY_PNG_BYTES)
    image_paragraph.add_run().add_picture(png_path)
    os.remove(png_path)
    document.save(path)
    print(f"Wrote {path}")


def generate_image_pptx():
    path = os.path.join(SAMPLES_DIR, "image-sample.pptx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank layout
    png_path = os.path.join(SAMPLES_DIR, "_tmp-image-fixture.png")
    with open(png_path, "wb") as f:
        f.write(_TINY_PNG_BYTES)
    slide.shapes.add_picture(png_path, left=0, top=0)
    os.remove(png_path)
    presentation.save(path)
    print(f"Wrote {path}")


def generate_password_protected_pdf():
    # docs/17-UNIT-TEST-PLAN.md's PDF Extraction plan requires a real
    # password-protected fixture (BR-003) - checked in rather than generated
    # inline by the C# test itself, same as every other sample here, since
    # building it needs pymupdf (Python), not anything available from a
    # C# xUnit test.
    path = os.path.join(SAMPLES_DIR, "password-protected-sample.pdf")
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "This content is behind a password.", fontsize=12)

    document.save(
        path,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="sample-password",
        owner_pw="sample-owner-password",
    )
    print(f"Wrote {path}")


if __name__ == "__main__":
    os.makedirs(SAMPLES_DIR, exist_ok=True)
    generate_txt()
    generate_docx()
    generate_xlsx()
    generate_pptx()
    generate_pdf()
    generate_image_pdf()
    generate_image_docx()
    generate_image_pptx()
    generate_password_protected_pdf()
