"""Generates a synthetic corpus for cost-testing the conversion credit ratios.

SR-BIL-4 defines a credit as 1 PDF page, 1 PPTX slide, 3,000 DOCX/TXT
characters, or 1,000 non-empty XLSX cells. Those baselines were chosen as a
deterministic, explainable unit - NOT from measured cost. This corpus exists so
they can be measured, because a credit that costs us ten times more to serve in
one format than another is a pricing bug waiting to be discovered by whichever
customer notices first.

The corpus deliberately spans a RANGE of sizes per format rather than one file
each. Cost per unit is only meaningful if it holds as documents grow: a
per-file fixed overhead that dominates a one-page PDF but vanishes on a
hundred-page one would make a single measurement actively misleading.

Synthetic and non-confidential (CLAUDE.md section 34). Real documents would
give better numbers, and that limitation is recorded with the results.

Run from the repo root:
    python scripts/generate-cost-corpus.py

Requires src/AI.Document.Converter.Python/requirements-dev.txt (pymupdf is used
for PDF authoring only - it is a development tool, never shipped).
"""

import os
import shutil

import pymupdf
import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CORPUS_DIR = os.path.join(REPO_ROOT, "artifacts", "cost-corpus")

# Ordinary prose rather than "lorem ipsum" repeated, so the extractors face
# realistic word and sentence boundaries. Repetition is unavoidable in a
# synthetic corpus and is noted as a limitation with the results.
PARAGRAPH = (
    "The quarterly review covers operating performance, outstanding risks and "
    "the actions agreed at the previous meeting. Figures are provisional until "
    "the audit completes. "
)


def text_of_length(target_characters):
    repeats = (target_characters // len(PARAGRAPH)) + 1
    return (PARAGRAPH * repeats)[:target_characters]


def make_pdf(path, pages):
    document = pymupdf.open()

    for page_number in range(pages):
        page = document.new_page()
        # Roughly a page of text, so a "page" of billable work is a realistic
        # amount of extraction rather than an almost-empty sheet.
        page.insert_textbox(
            pymupdf.Rect(50, 50, 545, 780),
            f"Page {page_number + 1}\n\n{text_of_length(2200)}",
            fontsize=10,
        )

    document.save(path)
    document.close()


def make_docx(path, characters):
    document = docx.Document()
    remaining = characters

    # Split across paragraphs, because one enormous paragraph is not what a
    # real document looks like and would understate paragraph-handling cost.
    while remaining > 0:
        chunk = min(remaining, 1500)
        document.add_paragraph(text_of_length(chunk))
        remaining -= chunk

    document.save(path)


def make_xlsx(path, cells):
    workbook = openpyxl.Workbook()
    sheet = workbook.active

    columns = 10
    rows = max(1, cells // columns)

    for row in range(1, rows + 1):
        for column in range(1, columns + 1):
            sheet.cell(row=row, column=column, value=f"r{row}c{column}")

    workbook.save(path)


def make_pptx(path, slides):
    presentation = Presentation()
    layout = presentation.slide_layouts[5]

    for slide_number in range(slides):
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = f"Slide {slide_number + 1}"

        box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(4))
        box.text_frame.text = text_of_length(600)

    presentation.save(path)


def make_txt(path, characters):
    with open(path, "w", encoding="utf-8") as handle:
        handle.write(text_of_length(characters))


# (filename, format, billable units, builder). Units are what SR-BIL-4 counts,
# so the credit cost of each file is derivable without re-running the pricing
# code - the measurement must not depend on the thing it is measuring.
CORPUS = [
    ("pdf-001p.pdf", "pdf", 1, lambda p: make_pdf(p, 1)),
    ("pdf-005p.pdf", "pdf", 5, lambda p: make_pdf(p, 5)),
    ("pdf-020p.pdf", "pdf", 20, lambda p: make_pdf(p, 20)),
    ("pdf-050p.pdf", "pdf", 50, lambda p: make_pdf(p, 50)),
    ("pdf-100p.pdf", "pdf", 100, lambda p: make_pdf(p, 100)),

    ("docx-003k.docx", "docx", 3_000, lambda p: make_docx(p, 3_000)),
    ("docx-015k.docx", "docx", 15_000, lambda p: make_docx(p, 15_000)),
    ("docx-060k.docx", "docx", 60_000, lambda p: make_docx(p, 60_000)),
    ("docx-150k.docx", "docx", 150_000, lambda p: make_docx(p, 150_000)),
    ("docx-300k.docx", "docx", 300_000, lambda p: make_docx(p, 300_000)),

    ("xlsx-001k.xlsx", "xlsx", 1_000, lambda p: make_xlsx(p, 1_000)),
    ("xlsx-005k.xlsx", "xlsx", 5_000, lambda p: make_xlsx(p, 5_000)),
    ("xlsx-020k.xlsx", "xlsx", 20_000, lambda p: make_xlsx(p, 20_000)),
    ("xlsx-050k.xlsx", "xlsx", 50_000, lambda p: make_xlsx(p, 50_000)),
    ("xlsx-100k.xlsx", "xlsx", 100_000, lambda p: make_xlsx(p, 100_000)),

    ("pptx-001s.pptx", "pptx", 1, lambda p: make_pptx(p, 1)),
    ("pptx-005s.pptx", "pptx", 5, lambda p: make_pptx(p, 5)),
    ("pptx-020s.pptx", "pptx", 20, lambda p: make_pptx(p, 20)),
    ("pptx-050s.pptx", "pptx", 50, lambda p: make_pptx(p, 50)),

    ("txt-003k.txt", "txt", 3_000, lambda p: make_txt(p, 3_000)),
    ("txt-015k.txt", "txt", 15_000, lambda p: make_txt(p, 15_000)),
    ("txt-060k.txt", "txt", 60_000, lambda p: make_txt(p, 60_000)),
    ("txt-150k.txt", "txt", 150_000, lambda p: make_txt(p, 150_000)),
    ("txt-300k.txt", "txt", 300_000, lambda p: make_txt(p, 300_000)),

    # Deliberately far larger than the rest. Plain text turned out to extract
    # so cheaply that at 300k characters the whole measurement sat inside the
    # engine's own startup noise - a number below the measurement floor is not
    # a number. These push it above.
    ("txt-1500k.txt", "txt", 1_500_000, lambda p: make_txt(p, 1_500_000)),
    ("txt-3000k.txt", "txt", 3_000_000, lambda p: make_txt(p, 3_000_000)),
    ("docx-1500k.docx", "docx", 1_500_000, lambda p: make_docx(p, 1_500_000)),
]


def main():
    if os.path.exists(CORPUS_DIR):
        shutil.rmtree(CORPUS_DIR)

    os.makedirs(CORPUS_DIR)

    for file_name, file_format, units, build in CORPUS:
        path = os.path.join(CORPUS_DIR, file_name)
        build(path)
        print(f"{file_name:20} {file_format:5} {units:>8} units  {os.path.getsize(path):>10} bytes")

    print(f"\n{len(CORPUS)} files in {CORPUS_DIR}")


if __name__ == "__main__":
    main()
